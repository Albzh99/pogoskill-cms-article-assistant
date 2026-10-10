import importlib.util
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from PIL import Image


SPEC = importlib.util.spec_from_file_location("mac_images", Path(__file__).with_name("cms-macos-image-pipeline.py"))
PIPE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(PIPE)


class MacImagePipelineTests(unittest.TestCase):
    def test_extract_accounts_for_repeated_docx_images(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            image = root / "source.png"
            Image.new("RGB", (40, 20), "blue").save(image)
            doc = Document()
            doc.add_picture(str(image))
            doc.add_paragraph("Between")
            doc.add_picture(str(image))
            source = root / "article.docx"
            doc.save(source)
            manifest = PIPE.extract(source, root / "images", "pikmin-guide")
            items = PIPE.read(manifest)["items"]
            self.assertEqual(["img-01", "img-02"], [item["image_key"] for item in items])
            self.assertEqual(items[0]["source_sha256"], items[1]["source_sha256"])
            self.assertNotEqual(items[0]["fallback_name"], items[1]["fallback_name"])

    def fixture(self, root):
        fallback = root / "pikmin-guide-01.png"
        Image.new("RGB", (40, 20), "red").save(fallback)
        item = {"image_key": "img-01", "fallback_path": str(fallback), "fallback_name": fallback.name,
                "webp_path": str(root / "pikmin-guide-01.webp"), "webp_name": "pikmin-guide-01.webp",
                "width": 40, "height": 20, "max_width": 40, "alt": "皮克敏圖片", "status": "extracted"}
        manifest = root / "image-manifest.json"
        PIPE.save(manifest, {"items": [item]})
        return manifest

    def test_conversion_preserves_fallback_and_dimensions(self):
        with tempfile.TemporaryDirectory() as temp:
            manifest = self.fixture(Path(temp))
            before = Path(PIPE.read(manifest)["items"][0]["fallback_path"]).read_bytes()
            PIPE.convert(manifest)
            item = PIPE.read(manifest)["items"][0]
            self.assertEqual("converted", item["status"])
            self.assertEqual(before, Path(item["fallback_path"]).read_bytes())
            with Image.open(item["webp_path"]) as image:
                self.assertEqual((40, 20), image.size)

    def test_upload_reuses_pair_without_duplicate_post(self):
        with tempfile.TemporaryDirectory() as temp:
            manifest = self.fixture(Path(temp))
            PIPE.convert(manifest)
            image = PIPE.read(manifest)["items"][0]
            row = lambda name: {"name": name, "w": 40, "h": 20,
                                "upload": "https://site.p.cms.afirstsoft.cn/example/" + name,
                                "online": "https://tw.pogoskill.com/images/pikmin/" + name}
            cms = Mock()
            cms.request_json.side_effect = [
                {"request_id": "dirs", "data": {}},
                {"request_id": "list", "data": {"list": [{"fileChildren": [row(image["fallback_name"]), row(image["webp_name"])]}]}}
            ]
            with patch.object(PIPE, "http_readable", return_value=True):
                PIPE.upload(manifest, 324, "pikmin", cms)
            self.assertEqual("reused_existing", PIPE.read(manifest)["items"][0]["status"])
            cms.cms_post.assert_not_called()

    def test_upload_new_pair_checkpoints_publish_evidence(self):
        with tempfile.TemporaryDirectory() as temp:
            manifest = self.fixture(Path(temp))
            PIPE.convert(manifest)
            item = PIPE.read(manifest)["items"][0]
            row = lambda name: {"name": name, "w": 40, "h": 20,
                                "upload": "https://site.p.cms.afirstsoft.cn/example/" + name,
                                "online": "https://tw.pogoskill.com/images/pikmin/" + name}
            rows = [row(item["fallback_name"]), row(item["webp_name"])]
            cms = Mock()
            cms.request_json.side_effect = [
                {"request_id": "dirs", "data": {}},
                {"request_id": "before", "data": {"list": [{"fileChildren": []}]}},
                {"request_id": "after", "data": {"list": [{"fileChildren": rows}]}}
            ]
            cms.multipart_image_body.return_value = (b"bytes", "multipart/form-data; boundary=ok")
            cms.cms_post.return_value = {"code": 0, "request_id": "upload-uuid", "data": {
                "publish_id": 99, "total": 2, "err_name_files": [], "list": rows}}
            with patch.object(PIPE, "http_readable", return_value=True):
                PIPE.upload(manifest, 324, "pikmin", cms)
            saved = PIPE.read(manifest)["items"][0]
            self.assertEqual("uploaded_pending_publish", saved["status"])
            self.assertEqual(99, saved["publish_id"])
            self.assertEqual("upload-uuid", PIPE.read(saved["upload_evidence_path"])["request_id"])
            self.assertIn("tw.pogoskill.com/images/pikmin/", saved["webp_public_url"])

    def test_publish_only_upload_evidence_ids_and_no_repeat(self):
        with tempfile.TemporaryDirectory() as temp:
            manifest = self.fixture(Path(temp))
            data = PIPE.read(manifest)
            data["items"][0].update(status="uploaded_pending_publish", publish_id=42,
                upload_request_id="upload-uuid", fallback_public_url="https://tw.pogoskill.com/images/pikmin/a.png",
                webp_public_url="https://tw.pogoskill.com/images/pikmin/a.webp")
            PIPE.save(manifest, data)
            cms = Mock()
            cms.request_json.return_value = {"request_id": "publish-uuid", "data": {"success": [{"id": 42}], "failed": []}}
            readable = Mock(side_effect=[False, True, True])
            PIPE.publish(manifest, cms, readable, lambda _: None)
            self.assertEqual("image_published", PIPE.read(manifest)["items"][0]["status"])
            self.assertEqual("/cms/pagepublish/publish", cms.request_json.call_args.args[0])
            self.assertEqual([42], cms.request_json.call_args.args[1]["ids"])
            PIPE.publish(manifest, cms, lambda _: True, lambda _: None)
            self.assertEqual(1, cms.request_json.call_count)

    def test_apply_uses_frontend_urls_and_escapes_ampersands(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            manifest = self.fixture(root)
            data = PIPE.read(manifest)
            item = data["items"][0]
            item.update(status="image_published", site_id=324,
                fallback_public_url="https://tw.pogoskill.com/images/pikmin/a.png?w=40&h=20",
                webp_public_url="https://tw.pogoskill.com/images/pikmin/a.webp?w=40&h=20")
            PIPE.save(manifest, data)
            source = root / "draft.html"
            source.write_text('<div class="img-wrap text-center"><!-- IMAGE_PENDING image-key: img-01 --></div>', encoding="utf-8")
            output = root / "ready.html"
            PIPE.apply(manifest, source, output)
            body = output.read_text(encoding="utf-8")
            self.assertIn("a.webp?w=40&amp;h=20", body)
            self.assertNotIn("site.p.cms.afirstsoft.cn", body)
            with self.assertRaises(ValueError):
                PIPE.apply(manifest, source, output)


if __name__ == "__main__":
    unittest.main()
